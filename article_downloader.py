"""
article_downloader.py
=====================
Modul ekstraksi DOI dari berbagai format dokumen (Excel, Word, PDF, TXT)
dan pengunduh otomatis artikel PDF ilmiah dengan strategi hybrid (Open Access / Unpaywall & Sci-Hub).
"""

import os
import re
import io
import time
import zipfile
import threading
from datetime import datetime
from typing import List, Dict, Any, Optional, Tuple

import requests
from bs4 import BeautifulSoup

# Direktori default penyimpanan artikel PDF
ARTIKEL_DIR = "artikel"

# Daftar mirror Sci-Hub resmi yang aktif
DEFAULT_SCIHUB_MIRRORS = [
    "https://sci-hub.ru",
    "https://sci-hub.su",
    "https://sci-hub.st",
    "https://sci-hub.red",
    "https://sci-hub.box"
]


def sanitize_filename(name: str, max_length: int = 80) -> str:
    """Membersihkan string agar aman dijadikan nama berkas di semua sistem operasi."""
    clean = re.sub(r'[\\/*?:"<>|]', '', str(name)).strip()
    clean = re.sub(r'\s+', '_', clean)
    clean = clean.strip('._')
    return clean[:max_length] if clean else "artikel"


def extract_raw_doi(doi_str: str) -> str:
    """Mengekstrak identifier murni 10.xxxx/... tanpa awalan URL atau trailing punctuation."""
    if not doi_str:
        return ""
    cleaned = re.sub(r'^(https?://(dx\.)?doi\.org/|doi:\s*|doi/)+', '', str(doi_str).strip(), flags=re.IGNORECASE)
    m = re.search(r'(10\.\d{4,9}/[^\s"\'<>]+)', cleaned)
    if m:
        doi = m.group(1)
        doi = re.sub(r'[\.,;:\)\]>]+$', '', doi)
        return doi
    return ""


def extract_dois_from_text(text: str) -> List[str]:
    """Mengekstrak seluruh DOI unik dari string teks bebas."""
    if not text:
        return []
    # Pola regex penangkapan DOI universal
    pattern = r'(?:https?://(?:dx\.)?doi\.org/|doi:\s*|doi/)?\b(10\.\d{4,9}/[^\s"\'<>]+)'
    matches = re.findall(pattern, text, flags=re.IGNORECASE)
    seen = set()
    results = []
    for m in matches:
        raw_doi = extract_raw_doi(m)
        if raw_doi and raw_doi.lower() not in seen:
            seen.add(raw_doi.lower())
            results.append(raw_doi)
    return results


def extract_dois_from_file(file_obj, filename: str) -> List[str]:
    """
    Mengekstrak daftar DOI dari berkas yang diunggah.
    Mendukung format: .xlsx, .xls, .docx, .pdf, .txt, .csv
    """
    ext = os.path.splitext(filename)[1].lower()
    full_text = []

    try:
        if ext in ['.xlsx', '.xls']:
            import pandas as pd
            xl = pd.ExcelFile(file_obj)
            for s in xl.sheet_names:
                df = pd.read_excel(xl, sheet_name=s)
                for col in df.columns:
                    full_text.append(str(col))
                    full_text.extend(df[col].dropna().astype(str).tolist())

        elif ext == '.docx':
            import docx
            doc = docx.Document(file_obj)
            for p in doc.paragraphs:
                if p.text:
                    full_text.append(p.text)
            for t in doc.tables:
                for row in t.rows:
                    for cell in row.cells:
                        if cell.text:
                            full_text.append(cell.text)

        elif ext == '.pdf':
            import pypdf
            reader = pypdf.PdfReader(file_obj)
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    full_text.append(t)

        elif ext in ['.txt', '.csv']:
            if isinstance(file_obj, (str, os.PathLike)):
                with open(file_obj, 'r', encoding='utf-8', errors='ignore') as f:
                    full_text.append(f.read())
            else:
                content = file_obj.read()
                if isinstance(content, bytes):
                    content = content.decode('utf-8', errors='ignore')
                full_text.append(content)
        else:
            # Fallback membaca sebagai teks mentah
            content = file_obj.read()
            if isinstance(content, bytes):
                content = content.decode('utf-8', errors='ignore')
            full_text.append(content)

    except Exception as err:
        print(f"[ERROR] Gagal memindai berkas {filename}: {err}")

    merged_text = "\n".join(full_text)
    return extract_dois_from_text(merged_text)


def fetch_crossref_title(doi: str, timeout: float = 4.0) -> str:
    """Mendapatkan judul publikasi dari Crossref API untuk penamaan file PDF yang rapi."""
    raw = extract_raw_doi(doi)
    if not raw:
        return ""
    try:
        url = f"https://api.crossref.org/works/{requests.utils.quote(raw)}"
        headers = {'User-Agent': 'PublikasiDosenScraper/1.0 (mailto:scholar_scraper@dosen.ac.id)'}
        r = requests.get(url, headers=headers, timeout=timeout)
        if r.status_code == 200:
            titles = r.json().get('message', {}).get('title', [])
            if titles:
                return titles[0]
    except Exception:
        pass
    return ""


def download_from_unpaywall(doi: str, output_path: str, timeout: float = 12.0) -> Tuple[bool, str]:
    """Mencoba mengunduh file PDF secara legal & cepat via Unpaywall Open Access API."""
    raw = extract_raw_doi(doi)
    if not raw:
        return False, "DOI tidak valid"
    try:
        api_url = f"https://api.unpaywall.org/v2/{requests.utils.quote(raw)}?email=scholar_scraper@dosen.ac.id"
        r = requests.get(api_url, timeout=5)
        if r.status_code == 200:
            data = r.json()
            if data.get('is_oa'):
                best_oa = data.get('best_oa_location') or {}
                pdf_url = best_oa.get('url_for_pdf') or best_oa.get('url')
                if pdf_url and ('pdf' in pdf_url.lower() or pdf_url.endswith('.pdf')):
                    headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'}
                    pdf_resp = requests.get(pdf_url, headers=headers, timeout=timeout, stream=True, verify=False)
                    if pdf_resp.status_code == 200:
                        first_chunk = next(pdf_resp.iter_content(1024), b'')
                        if first_chunk.startswith(b'%PDF') or 'application/pdf' in pdf_resp.headers.get('Content-Type', '').lower():
                            with open(output_path, 'wb') as f:
                                f.write(first_chunk)
                                for chunk in pdf_resp.iter_content(8192):
                                    if chunk:
                                        f.write(chunk)
                            return True, "Open Access (Legal)"
    except Exception:
        pass
    return False, "Tidak tersedia di Open Access"


def download_from_scihub(doi: str, output_path: str, mirrors: Optional[List[str]] = None, timeout: float = 15.0) -> Tuple[bool, str]:
    """Mencoba mengunduh artikel PDF via jaringan mirror Sci-Hub."""
    raw = extract_raw_doi(doi)
    if not raw:
        return False, "DOI tidak valid"

    active_mirrors = mirrors or DEFAULT_SCIHUB_MIRRORS
    headers = {
        'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
    }

    for mirror in active_mirrors:
        base_mirror = mirror.rstrip("/")
        page_url = f"{base_mirror}/{raw}"
        try:
            r = requests.get(page_url, headers=headers, timeout=7, verify=False)
            if r.status_code != 200 or 'altcha' in r.text.lower() or 'robot' in r.text.lower():
                continue

            soup = BeautifulSoup(r.text, 'html.parser')
            pdf_url = ""

            # 1. Cek elemen object, embed, iframe
            for tag in ['object', 'embed', 'iframe']:
                el = soup.find(tag)
                if el:
                    src = el.get('data') or el.get('src')
                    if src and ('.pdf' in src or '/storage/' in src or 'download' in src):
                        pdf_url = src
                        break

            # 2. Cek button location.href
            if not pdf_url:
                btn = soup.find('button', onclick=True)
                if btn:
                    m = re.search(r"location\.href=['\"]([^'\"]+)['\"]", btn['onclick'])
                    if m:
                        pdf_url = m.group(1)

            # 3. Cek tag <a>
            if not pdf_url:
                for a in soup.find_all('a', href=True):
                    href = a['href']
                    if '.pdf' in href or '/storage/' in href:
                        pdf_url = href
                        break

            if not pdf_url:
                continue

            # Bersihkan fragment #...
            pdf_url = pdf_url.split('#')[0]

            # Normalisasi URL
            if pdf_url.startswith('//'):
                pdf_url = 'https:' + pdf_url
            elif pdf_url.startswith('/'):
                pdf_url = base_mirror + pdf_url
            elif not pdf_url.startswith('http'):
                pdf_url = f"{base_mirror}/{pdf_url}"

            # Unduh file PDF
            req_headers = dict(headers)
            req_headers['Referer'] = page_url
            pdf_resp = requests.get(pdf_url, headers=req_headers, timeout=timeout, verify=False, stream=True)
            if pdf_resp.status_code == 200:
                first_chunk = next(pdf_resp.iter_content(1024), b'')
                if first_chunk.startswith(b'%PDF') or 'application/pdf' in pdf_resp.headers.get('Content-Type', '').lower():
                    with open(output_path, 'wb') as f:
                        f.write(first_chunk)
                        for chunk in pdf_resp.iter_content(8192):
                            if chunk:
                                f.write(chunk)
                    return True, f"Sci-Hub ({base_mirror.replace('https://', '')})"
        except Exception:
            continue

    return False, "Tidak ditemukan di seluruh mirror Sci-Hub"


class ArticleDownloaderJob:
    """Manajer tugas latar belakang untuk pengunduhan daftar DOI artikel ilmiah."""

    def __init__(
        self,
        doi_list: List[str],
        subfolder_name: str,
        mirrors: Optional[List[str]] = None,
        delay_sec: float = 2.0
    ):
        self.doi_list = [extract_raw_doi(d) for d in doi_list if extract_raw_doi(d)]
        # Hilangkan duplikat berurutan
        seen = set()
        unique_dois = []
        for d in self.doi_list:
            if d.lower() not in seen:
                seen.add(d.lower())
                unique_dois.append(d)
        self.doi_list = unique_dois

        self.clean_folder_name = sanitize_filename(subfolder_name) or f"unduhan_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        self.output_dir = os.path.join(ARTIKEL_DIR, self.clean_folder_name)
        os.makedirs(self.output_dir, exist_ok=True)

        self.mirrors = mirrors or DEFAULT_SCIHUB_MIRRORS
        self.delay_sec = max(0.5, float(delay_sec))

        self.is_running = False
        self.stop_requested = False
        self.current_idx = 0
        self.total_items = len(self.doi_list)
        self.current_doi = ""
        self.status = "Menunggu..."
        self.logs: List[str] = []
        self.results: List[Dict[str, Any]] = []
        self.thread: Optional[threading.Thread] = None

    def add_log(self, msg: str):
        ts = datetime.now().strftime("%H:%M:%S")
        self.logs.append(f"[{ts}] {msg}")
        self.logs = self.logs[-200:]

    def _worker(self):
        self.is_running = True
        self.stop_requested = False
        self.results = []
        self.current_idx = 0
        self.add_log(f"Memulai tugas pengunduhan {self.total_items} artikel ke folder '{self.output_dir}'.")

        try:
            for idx, raw_doi in enumerate(self.doi_list, 1):
                if self.stop_requested:
                    self.add_log("Pengunduhan dihentikan oleh pengguna.")
                    break

                self.current_idx = idx
                self.current_doi = raw_doi
                self.status = f"Memproses ({idx}/{self.total_items}): {raw_doi}..."
                self.add_log(f"[{idx}/{self.total_items}] Mencari artikel: {raw_doi}")

                # Cari judul publikasi via Crossref API
                title = fetch_crossref_title(raw_doi) or ""
                safe_title = sanitize_filename(title, max_length=60) if title else ""
                clean_doi_part = sanitize_filename(raw_doi.replace('/', '_'), max_length=40)
                filename = f"{safe_title}_{clean_doi_part}.pdf" if safe_title else f"{clean_doi_part}.pdf"
                output_pdf_path = os.path.join(self.output_dir, filename)

                success = False
                source = "-"

                # 1. Coba Open Access terlebih dahulu (Cepat & Legal)
                if not success and not self.stop_requested:
                    ok_oa, src_oa = download_from_unpaywall(raw_doi, output_pdf_path)
                    if ok_oa:
                        success = True
                        source = src_oa

                # 2. Jika bukan Open Access, coba Sci-Hub multi-mirror
                if not success and not self.stop_requested:
                    ok_sh, src_sh = download_from_scihub(raw_doi, output_pdf_path, mirrors=self.mirrors)
                    if ok_sh:
                        success = True
                        source = src_sh

                # Catat hasil
                if success and os.path.exists(output_pdf_path):
                    f_size_kb = round(os.path.getsize(output_pdf_path) / 1024, 1)
                    res_item = {
                        'No': idx,
                        'DOI': f"https://doi.org/{raw_doi}",
                        'Judul Artikel': title or "-",
                        'Sumber': source,
                        'Status': "Berhasil",
                        'Ukuran': f"{f_size_kb} KB",
                        'Nama File': filename,
                        'Path': output_pdf_path
                    }
                    self.add_log(f"[{idx}/{self.total_items}] BERHASIL diunduh via {source} ({f_size_kb} KB).")
                else:
                    res_item = {
                        'No': idx,
                        'DOI': f"https://doi.org/{raw_doi}",
                        'Judul Artikel': title or "-",
                        'Sumber': "-",
                        'Status': "Tidak Ditemukan",
                        'Ukuran': "-",
                        'Nama File': "-",
                        'Path': ""
                    }
                    self.add_log(f"[{idx}/{self.total_items}] TIDAK DITEMUKAN di Open Access maupun Sci-Hub.")

                self.results.append(res_item)

                # Jeda sopan antar unduhan
                if idx < self.total_items and not self.stop_requested:
                    time.sleep(self.delay_sec)

        except Exception as e:
            self.add_log(f"Error selama proses unduh: {e}")
            self.status = f"Terjadi kesalahan: {e}"
        finally:
            self.is_running = False
            succ_count = sum(1 for r in self.results if r['Status'] == 'Berhasil')
            if self.stop_requested:
                self.status = f"Dihentikan. {succ_count}/{len(self.results)} artikel berhasil diunduh."
            else:
                self.status = f"Selesai! {succ_count}/{self.total_items} artikel berhasil disimpan di folder '{self.clean_folder_name}'."
            self.add_log(self.status)

    def start(self):
        if self.is_running:
            return
        self.thread = threading.Thread(target=self._worker, daemon=True)
        self.thread.start()

    def stop(self):
        self.stop_requested = True
        self.add_log("Permintaan stop pengunduhan diterima...")
